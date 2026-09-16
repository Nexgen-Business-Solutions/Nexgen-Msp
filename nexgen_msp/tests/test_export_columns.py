"""What a sheet takes away: the columns that were picked, and one block per service."""

import io

import frappe
from openpyxl import load_workbook

from nexgen_msp.api.internal.services.service_lifecycle_service import ServiceLifecycleService
from nexgen_msp.utils import export_columns

from .base import MSPTestCase

CATALOGUE = [
    ("full_name", "User"),
    ("username", "Username"),
    ("department", "Department"),
]


class TestTheColumnsThatWerePicked(MSPTestCase):
    def test_only_what_was_picked_comes_out_in_the_catalogue_order(self):
        picked = export_columns.chosen(CATALOGUE, ["department", "full_name", "nothing_like_it"])

        self.assertEqual(picked, [("full_name", "User"), ("department", "Department")])

    def test_picking_nothing_leaves_the_sheet_as_it_has_always_been(self):
        self.assertEqual(export_columns.chosen(CATALOGUE, None), CATALOGUE)
        self.assertEqual(export_columns.chosen(CATALOGUE, "[]"), CATALOGUE)

    def test_a_selection_travels_as_json(self):
        self.assertEqual(export_columns.chosen(CATALOGUE, '["username"]'), [("username", "Username")])

    def test_a_selection_of_keys_nobody_knows_is_not_an_empty_sheet(self):
        self.assertEqual(export_columns.chosen(CATALOGUE, ["gone", "older"]), CATALOGUE)


class TestOneBlockPerService(MSPTestCase):
    def setUp(self):
        super().setUp()
        self.tag = frappe.generate_hash(length=6)
        self.customer = self.make_customer(self.tag)
        self.track("MSP Approval Authority", self.customer)
        self.john = self.make_person(self.customer, "John")
        self.first = self.offering("EX1")
        self.second = self.offering("EX2")
        self.laptop = self.make_device(
            self.customer, hostname=f"EXP{self.tag[:3]}", holder=self.john, serial=f"EX-{self.tag}"
        )
        self.on_machine = self.offering("EX3", scope="Device")

        self.open_on(self.first, client_user=self.john, on=frappe.utils.add_days(frappe.utils.today(), -40))
        self.open_on(self.second, client_user=self.john, on=frappe.utils.add_days(frappe.utils.today(), -10))
        self.open_on(self.on_machine, managed_device=self.laptop, scope="Device")

    def offering(self, suffix, scope="User"):
        service = self.make_service(f"{suffix}{self.tag[:3]}", scope=scope)
        self.cover_service(self.customer, service)

        return service

    def open_on(self, service, scope="User", on=None, **target):
        out = ServiceLifecycleService.activate(
            customer=self.customer,
            service_item=service,
            target_scope=scope,
            effective_date=on,
            **target,
        )

        return self.track("MSP Service Assignment", out["name"])

    def test_each_service_gets_its_own_consecutive_columns(self):
        rows = [{"name": self.john}]
        held = export_columns.of_people([self.john])

        columns = export_columns.service_columns(
            rows, held, ["service_name", "effective_start_date", "last_billed_on"]
        )

        self.assertEqual(
            [key for key, _label in columns][:3],
            ["service_1_service_name", "service_1_effective_start_date", "service_1_last_billed_on"],
        )
        self.assertEqual(
            [label for _key, label in columns][:3],
            ["Service 1", "Service 1 · Since", "Service 1 · Last billed on"],
        )
        # their machine's service is theirs to read too: three services, three blocks
        self.assertEqual(len(columns), 9)
        self.assertEqual(rows[0]["service_1_service_name"], frappe.db.get_value("Item", self.first, "item_name"))
        self.assertEqual(
            str(rows[0]["service_1_effective_start_date"]),
            str(frappe.utils.add_days(frappe.utils.today(), -40)),
        )

    def test_a_machine_reads_the_services_running_on_it(self):
        rows = [{"name": self.laptop}]
        running = export_columns.of_devices([self.laptop])

        columns = export_columns.service_columns(rows, running, ["service_name", "holder_name"])

        self.assertEqual([key for key, _label in columns], ["service_1_service_name", "service_1_holder_name"])
        self.assertEqual(rows[0]["service_1_holder_name"], frappe.db.get_value("MSP Client User", self.john, "full_name"))

    def test_a_sheet_that_asked_for_no_service_carries_none(self):
        rows = [{"name": self.john}]

        self.assertEqual(export_columns.service_columns(rows, export_columns.of_people([self.john]), []), [])
        self.assertNotIn("service_1_service_name", rows[0])

    def test_the_file_comes_out_with_the_picked_headers(self):
        from nexgen_msp.api.internal.endpoints import v1

        v1.export_users(
            customer=self.customer,
            columns=["full_name", "department"],
            service_columns=["service_name", "effective_start_date"],
        )

        sheet = load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
        headers = [cell.value for cell in sheet[1]]

        self.assertEqual(headers[:2], ["User", "Department"])
        self.assertEqual(headers[2:4], ["Service 1", "Service 1 · Since"])
        self.assertNotIn("Username", headers)


class TestWhatACustomerTakesAway(MSPTestCase):
    """Their own people and machines, read as fully as we read them — bar our own notes."""

    def setUp(self):
        super().setUp()
        self.tag = frappe.generate_hash(length=6)
        self.customer = self.make_customer(self.tag)
        self.track("MSP Approval Authority", self.customer)
        self.john = self.make_person(self.customer, "John")
        self.laptop = self.make_device(
            self.customer,
            hostname=f"CUS{self.tag[:3]}",
            holder=self.john,
            serial=f"CU-{self.tag}",
        )
        frappe.db.set_value(
            "MSP Managed Device", self.laptop, {"model": "ThinkPad T14", "operating_system": "Windows 11"}
        )
        frappe.db.commit()
        self.manager = self.make_account("customer", "MSP Customer Manager", self.customer, suffix=f"ex{self.tag[:3]}")

    def as_manager(self, fn):
        frappe.set_user(self.manager)
        try:
            return fn()
        finally:
            frappe.set_user("Administrator")

    def headers(self):
        sheet = load_workbook(io.BytesIO(frappe.local.response.filecontent)).active

        return [cell.value for cell in sheet[1]]

    def row(self, header):
        sheet = load_workbook(io.BytesIO(frappe.local.response.filecontent)).active
        titles = [cell.value for cell in sheet[1]]

        return sheet.cell(row=2, column=titles.index(header) + 1).value

    def test_a_customer_reads_the_serials_and_types_of_the_machines_they_hold(self):
        from nexgen_msp.api.portal.endpoints import v1

        self.as_manager(
            lambda: v1.export_client_users(
                columns=["full_name", "serial_numbers", "device_type", "hostnames", "current_devices"]
            )
        )

        self.assertEqual(self.headers(), ["Name", "Devices", "Serial numbers", "Device type", "Devices held"])
        self.assertEqual(self.row("Serial numbers"), f"CU-{self.tag}")
        self.assertEqual(self.row("Devices held"), 1)

    def test_a_machine_sheet_carries_what_the_machine_is(self):
        from nexgen_msp.api.portal.endpoints import v1

        self.as_manager(
            lambda: v1.export_devices(columns=["hostname", "model", "operating_system", "previous_holders"])
        )

        self.assertEqual(self.row("Model"), "ThinkPad T14")
        self.assertEqual(self.row("Operating system"), "Windows 11")
        self.assertIn("John", self.row("Previous holders") or "")

    def test_our_own_notes_and_their_own_name_are_not_on_offer(self):
        from nexgen_msp.api.portal.endpoints import v1

        offered = {key for key, _label in v1.EXPORT_COLUMNS["users"]}
        offered |= {key for key, _label in v1.EXPORT_COLUMNS["devices"]}

        self.assertNotIn("remarks", offered)
        self.assertNotIn("customer", offered)
